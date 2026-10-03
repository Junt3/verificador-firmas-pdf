import os
import tkinter as tk
from tkinter import filedialog, messagebox
from pypdf import PdfReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def formatear_fecha(fecha_pdf):
    if not fecha_pdf or not isinstance(fecha_pdf, str) or not fecha_pdf.startswith("D:"):
        return fecha_pdf
    try:
        raw = fecha_pdf[2:16]
        if len(raw) == 14:
            return f"{raw[0:4]}-{raw[4:6]}-{raw[6:8]} {raw[8:10]}:{raw[10:12]}:{raw[12:14]}"
    except Exception:
        pass
    return fecha_pdf

def analizar_firmas_pdf(ruta_archivo):
    try:
        reader = PdfReader(ruta_archivo)
        fields = reader.get_fields()
        is_signed = False
        nombres, fechas = [], []

        if fields:
            for key, field in fields.items():
                if field.get('/FT') == '/Sig':
                    is_signed = True
                    sig_data = field.get('/V')
                    if sig_data:
                        nombre = sig_data.get('/Name', 'No especificado')
                        fecha = sig_data.get('/M', 'No especificada')
                        nombres.append(str(nombre))
                        fechas.append(formatear_fecha(str(fecha)))

        return ("Sí" if is_signed else "No", " | ".join(nombres) if nombres else "-", " | ".join(fechas) if fechas else "-")
    except Exception:
        return ("Error", "Archivo protegido o corrupto", "-")

def main():
    root = tk.Tk()
    root.withdraw()

    # 1. Selección de carpeta
    archivo_guia = filedialog.askopenfilename(
        title="Selecciona un PDF de la carpeta a analizar",
        filetypes=[("Archivos PDF", "*.pdf")]
    )
    
    if not archivo_guia:
        return

    carpeta = os.path.dirname(archivo_guia)
    archivos_pdf = [f for f in os.listdir(carpeta) if f.lower().endswith('.pdf')]

    if not archivos_pdf:
        messagebox.showwarning("Sin archivos", "No se encontraron PDFs en la carpeta.")
        return

    # 2. Selección de destino para el Excel
    ruta_excel = filedialog.asksaveasfilename(
        title="¿Dónde quieres guardar el reporte de Excel?",
        defaultextension=".xlsx",
        filetypes=[("Libro de Excel", "*.xlsx")],
        initialfile="Reporte_Firmas_Electronicas.xlsx"
    )

    if not ruta_excel:
        print("Operación cancelada por el usuario.")
        return

    # 3. Creación del Excel y Estilos (VERSIÓN A PRUEBA DE FALLOS)
    wb = Workbook()
    ws = wb.active
    ws.title = "Control de Firmas"

    # Letra negra en negrita (Garantiza visibilidad siempre)
    header_font = Font(bold=True, color="000000") 
    
    # Fondo sólido usando fgColor (Mucho más estable en openpyxl)
    header_fill = PatternFill(start_color="FFADD8E6", end_color="FFADD8E6", fill_type="solid")
    
    center_aligned = Alignment(horizontal="center", vertical="center")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                    top=Side(style='thin'), bottom=Side(style='thin'))

    # Encabezados
    headers = ["Nombre del Archivo", "¿Está Firmado?", "Fecha de la Firma", "Nombre del Firmante"]
    ws.append(headers)

    # Aplicar estilos explícitamente a cada celda de la primera fila
    for row in ws.iter_rows(min_row=1, max_row=1, min_col=1, max_col=4):
        for cell in row:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_aligned
            cell.border = border

    # 4. Procesar archivos y llenar datos
    print(f"Procesando {len(archivos_pdf)} archivos...")
    for nombre_archivo in archivos_pdf:
        ruta_completa = os.path.join(carpeta, nombre_archivo)
        firmado, nombres, fechas = analizar_firmas_pdf(ruta_completa)
        
        row = [nombre_archivo, firmado, fechas, nombres]
        ws.append(row)
        
        # Aplicar bordes al contenido generado
        for cell in ws[ws.max_row]:
            cell.border = border

    # 5. Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column].width = max_length + 2

    # Guardar archivo
    try:
        wb.save(ruta_excel)
        messagebox.showinfo("Éxito", f"Reporte generado correctamente en:\n{ruta_excel}")
        print(f"Proceso finalizado. Excel guardado en: {ruta_excel}")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo guardar el archivo: {e}")

if __name__ == "__main__":
    main()